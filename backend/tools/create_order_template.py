from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


columns = [
    "external_order_id", "sku", "product_description", "quantity", "unit_price",
    "currency", "supplier_code", "requested_by", "requested_at", "required_procedure_code",
]
target = Path(__file__).resolve().parents[1] / "templates" / "order_import_columns.xlsx"
workbook = Workbook()
sheet = workbook.active
sheet.title = "Pedidos"
sheet.append(columns)
sheet.append(["PED-PLANTILLA-001", "SKU-EJEMPLO", "Descripción del producto", 10, 12.5, "EUR", "SALTOKI", "laura.martin@smartwarehouse.local", "2026-09-25 10:00:00", "PURCHASE_APPROVAL"])
for cell in sheet[1]:
    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="2563EB")
    cell.alignment = Alignment(horizontal="center")
for column, width in zip("ABCDEFGHIJ", [24, 18, 34, 12, 14, 12, 18, 34, 24, 28]):
    sheet.column_dimensions[column].width = width
sheet.freeze_panes = "A2"
sheet.auto_filter.ref = f"A1:J2"
workbook.save(target)
print(target)
