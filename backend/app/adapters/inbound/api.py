import csv
import io
import json
import re
import shutil
import subprocess
import tempfile
from typing import Annotated, Callable
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from pypdf import PdfReader
from pydantic import BaseModel, Field

from ...application.services import WarehouseService
from ...application.ai_service import AIService
from ...application.integrations import IntegrationService
from ...application.auth import create_token, decode_token, verify_password
from ...domain.errors import OrderNotFoundError


class ValidationResponse(BaseModel):
    status: str
    risk: str
    reasons: list[str]


class HealthResponse(BaseModel):
    status: str
    service: str
    environment: str


class CreateOrderRequest(BaseModel):
    sku: str
    quantity: float = Field(gt=0)
    unit_price: float = Field(gt=0)
    supplier_code: str
    warehouse_code: str = "MAD-01"
    requester_email: str = "laura.martin@smartwarehouse.local"
    external_order_id: str | None = None


class AccountingExportRequest(BaseModel):
    target_system: str = "corporate-accounting-rest"


class ReceiptLineRequest(BaseModel):
    sku: str = Field(min_length=1)
    received_quantity: float = Field(ge=0)
    damaged_quantity: float = Field(default=0, ge=0)
    damage_reason: str | None = None


class CreateReceiptRequest(BaseModel):
    order_external_id: str
    receipt_number: str | None = None
    dock_code: str | None = None
    received_at: str | None = None
    lines: list[ReceiptLineRequest] = Field(min_length=1)


class StockMovementRequest(BaseModel):
    sku: str
    movement_type: str = Field(pattern="^(entry|exit|reserve|release|adjustment)$")
    quantity: float = Field(default=1, gt=0)
    adjustment_quantity: float | None = None
    warehouse_code: str = "MAD-01"
    reference_type: str | None = None
    reference_id: str | None = None
    reason: str | None = None


class OrderStatusRequest(BaseModel):
    status: str = Field(pattern="^(pending|validated|approved|sent_to_supplier|received|closed)$")
    reason: str | None = None


class AlertRuleRequest(BaseModel):
    code: str = Field(min_length=3, max_length=64)
    name: str = Field(min_length=3, max_length=160)
    description: str | None = None
    event_type: str = Field(min_length=3, max_length=64)
    severity: str = Field(pattern="^(info|warning|critical)$")
    enabled: bool = True
    recipients: list[str] = Field(default_factory=list)
    channels: list[str] = Field(default_factory=lambda: ["in_app"])


class AIChatRequest(BaseModel):
    message: str = Field(min_length=2, max_length=2000)


class LoginRequest(BaseModel):
    email: str = Field(min_length=5, max_length=190)
    password: str = Field(min_length=4, max_length=128)


class IntegrationSyncRequest(BaseModel):
    direction: str = Field(default="outbound", pattern="^(inbound|outbound|bidirectional)$")


def _amount(value: str) -> float | None:
    normalized = value.replace("€", "").replace("EUR", "").replace(" ", "").strip()
    if not normalized:
        return None


def _parse_order_import(filename: str, content: bytes) -> tuple[list[dict], list[dict]]:
    if filename.endswith(".csv"):
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        source_headers = reader.fieldnames or []
        rows = list(reader)
    elif filename.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        values = list(workbook.active.values)
        source_headers = [str(value).strip() if value is not None else "" for value in values[0]] if values else []
        rows = [dict(zip(source_headers, row)) for row in values[1:]]
    else:
        raise HTTPException(status_code=415, detail="Solo se admiten archivos .csv o .xlsx")
    aliases = {"ref": "sku", "sku": "sku", "producto": "product_description", "descripción": "product_description", "descripcion": "product_description", "cant.": "quantity", "cantidad": "quantity", "unidades": "quantity", "precio": "unit_price", "precio unitario": "unit_price", "proveedor": "supplier_code", "supplier_code": "supplier_code", "external_order_id": "external_order_id"}
    headers = {aliases.get(str(key).strip().lower(), str(key).strip()) for key in source_headers}
    required = {"sku", "quantity", "unit_price", "supplier_code"}
    errors = []
    if not rows:
        errors.append({"row": 1, "field": "headers", "message": "El archivo no contiene filas"})
    for field in sorted(required - headers):
        errors.append({"row": 1, "field": field, "message": f"Falta la columna obligatoria: {field}"})
    normalized = []
    for index, row in enumerate(rows, start=2):
        mapped = {aliases.get(str(key).strip().lower(), str(key).strip()): value for key, value in row.items()}
        if not any(value not in (None, "") for value in mapped.values()):
            continue
        mapped["_row"] = index
        for field in ("sku", "supplier_code"):
            if not mapped.get(field):
                errors.append({"row": index, "field": field, "message": f"Campo obligatorio vacío: {field}"})
        try:
            mapped["quantity"] = float(mapped.get("quantity", 0))
            if mapped["quantity"] <= 0:
                raise ValueError
        except (TypeError, ValueError):
            errors.append({"row": index, "field": "quantity", "message": "La cantidad debe ser numérica y mayor que cero"})
        try:
            mapped["unit_price"] = float(mapped.get("unit_price", 0))
            if mapped["unit_price"] <= 0:
                raise ValueError
        except (TypeError, ValueError):
            errors.append({"row": index, "field": "unit_price", "message": "El precio debe ser numérico y mayor que cero"})
        mapped.setdefault("supplier_code", "SALTOKI")
        normalized.append(mapped)
    return normalized, errors
    if "," in normalized and "." in normalized:
        normalized = normalized.replace(".", "").replace(",", ".")
    else:
        normalized = normalized.replace(",", ".")
    try:
        return round(float(normalized), 2)
    except ValueError:
        return None


def _accounting_file(rows: list[dict], export_format: str) -> tuple[bytes, str, str]:
    columns = [
        ("invoice_number", "invoice_number"), ("invoice_date", "invoice_date"),
        ("supplier_code", "supplier_code"), ("supplier", "supplier"),
        ("supplier_tax_id", "supplier_tax_id"), ("currency", "currency"),
        ("subtotal", "subtotal"), ("tax_amount", "tax_amount"), ("total", "total"),
        ("accounting_status", "accounting_status"), ("reconciliation_status", "reconciliation_status"),
        ("order_number", "order_number"), ("receipt_number", "receipt_number"),
    ]
    labels = [label for label, _ in columns]
    if export_format == "csv":
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=labels, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({label: row.get(key) for label, key in columns})
        return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "smart-warehouse-accounting.csv"
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Facturas"
    sheet.append(labels)
    for row in rows:
        sheet.append([row.get(key) for _, key in columns])
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "smart-warehouse-accounting.xlsx"


def _events_file(rows: list[dict], export_format: str) -> tuple[bytes, str, str]:
    columns = ["id", "created_at", "event_type", "severity", "aggregate_type", "aggregate_id", "actor_type", "actor_id", "payload"]
    normalized = []
    for row in rows:
        item = dict(row)
        payload = item.get("payload")
        item["payload"] = json.dumps(payload, ensure_ascii=False) if isinstance(payload, (dict, list)) else payload
        normalized.append(item)
    if export_format == "csv":
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(normalized)
        return output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8", "smart-warehouse-events.csv"
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Eventos"
    sheet.append(columns)
    for row in normalized:
        sheet.append([row.get(column) for column in columns])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "smart-warehouse-events.xlsx"


def _ocr_pdf(content: bytes) -> tuple[str, dict]:
    pdftoppm = shutil.which("pdftoppm")
    tesseract = shutil.which("tesseract")
    metadata = {"ocr_attempted": True, "ocr_available": bool(pdftoppm and tesseract), "ocr_pages": 0}
    if not pdftoppm or not tesseract:
        metadata["ocr_reason"] = "OCR no disponible: se necesita pdftoppm y tesseract"
        return "", metadata
    with tempfile.TemporaryDirectory(prefix="smart-warehouse-ocr-") as directory:
        pdf_path = Path(directory) / "invoice.pdf"
        prefix = Path(directory) / "page"
        pdf_path.write_bytes(content)
        try:
            subprocess.run([pdftoppm, "-png", "-r", "200", str(pdf_path), str(prefix)], check=True, capture_output=True, timeout=60)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
            metadata["ocr_reason"] = f"No se pudo rasterizar el PDF: {error}"
            return "", metadata
        pages = sorted(Path(directory).glob("page-*.png"))
        metadata["ocr_pages"] = len(pages)
        text_parts = []
        for page in pages:
            try:
                result = subprocess.run([tesseract, str(page), "stdout", "-l", "spa+eng", "--psm", "6"], check=True, capture_output=True, text=True, timeout=60)
                text_parts.append(result.stdout)
            except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
                continue
        text = "\n".join(text_parts).strip()
        if not text:
            metadata["ocr_reason"] = "OCR ejecutado pero no detectó texto legible"
        return text, metadata


def _extract_invoice(content: bytes) -> dict:
    try:
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception as error:
        raise HTTPException(status_code=422, detail=f"No se pudo leer el PDF: {error}") from error

    extraction_source = "digital" if text else "ocr"
    ocr_metadata = {"ocr_attempted": False, "ocr_available": False, "ocr_pages": 0}
    if not text:
        text, ocr_metadata = _ocr_pdf(content)
    invoice_match = re.search(r"(?:factura|invoice)\s*(?:n[ºo°.]?|number|no\.?)?\s*[:#-]?\s*([A-Z0-9][A-Z0-9/_-]+)", text, re.IGNORECASE)
    date_match = re.search(r"(?:fecha|date)\s*[:#-]?\s*(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})", text, re.IGNORECASE)
    tax_id_match = re.search(r"\b([A-Z]\d{8}|\d{8}[A-Z])\b", text.upper())
    total_match = re.search(r"(?:total(?:\s+factura)?|importe\s+total)\s*[:#-]?\s*([0-9][0-9., ]*)", text, re.IGNORECASE)
    subtotal_match = re.search(r"(?:subtotal|base\s+imponible)\s*[:#-]?\s*([0-9][0-9., ]*)", text, re.IGNORECASE)
    tax_match = re.search(r"(?:iva|impuestos?|tax)\s*[:#-]?\s*([0-9][0-9., ]*)", text, re.IGNORECASE)
    invoice_date = None
    if date_match:
        day, month, year = re.split(r"[/-]", date_match.group(1))
        invoice_date = f"{int(year):04d}-{int(month):02d}-{int(day):02d}" if len(year) == 4 else f"20{int(year):02d}-{int(month):02d}-{int(day):02d}"
    required = ("supplier_tax_id", "invoice_number", "invoice_date", "total")
    values = {
        "supplier_tax_id": tax_id_match.group(1) if tax_id_match else None,
        "invoice_number": invoice_match.group(1).upper() if invoice_match else None,
        "invoice_date": invoice_date,
        "currency": "EUR",
        "subtotal": _amount(subtotal_match.group(1)) if subtotal_match else None,
        "tax": _amount(tax_match.group(1)) if tax_match else None,
        "total": _amount(total_match.group(1)) if total_match else None,
        "pages": len(reader.pages),
        "text_detected": bool(text.strip()),
    }
    base_confidence = 0.78 if extraction_source == "ocr" else 0.96
    low_confidence_fields = [field for field in required if values.get(field) in (None, "")]
    field_confidence = {field: 0.0 if values.get(field) in (None, "") else base_confidence for field in required}
    complete = bool(text.strip()) and not low_confidence_fields
    review_reasons = []
    if extraction_source == "ocr":
        review_reasons.append("Datos obtenidos mediante OCR; requieren revisión humana")
    if low_confidence_fields:
        review_reasons.append(f"Campos no identificados: {', '.join(low_confidence_fields)}")
    extracted = {**values, **ocr_metadata, "extraction_source": extraction_source, "field_confidence": field_confidence, "low_confidence_fields": low_confidence_fields, "review_reasons": review_reasons}
    extracted["extraction_status"] = "extracted" if complete and extraction_source == "digital" else "needs_review"
    extracted["confidence"] = base_confidence if complete else (0.45 if text.strip() and extraction_source == "digital" else 0.35 if text.strip() else 0.0)
    return extracted


def build_router(service_provider: Callable[[], WarehouseService], environment: str, ai_service_provider: Callable[[], AIService] | None = None, auth_repository_provider: Callable[[], object] | None = None, auth_secret: str = "local-only-change-before-production", auth_token_ttl_seconds: int = 28800, integration_service_provider: Callable[[], IntegrationService] | None = None) -> APIRouter:
    router = APIRouter()

    def service() -> WarehouseService:
        return service_provider()

    def ai_service() -> AIService:
        if ai_service_provider is None:
            raise HTTPException(status_code=503, detail="Proveedor de IA no configurado")
        return ai_service_provider()

    def integration_service() -> IntegrationService:
        if integration_service_provider is None:
            raise HTTPException(status_code=503, detail="Integraciones no configuradas")
        return integration_service_provider()

    @router.post("/auth/login", tags=["auth"])
    def login(request: LoginRequest):
        if auth_repository_provider is None:
            raise HTTPException(status_code=503, detail="Autenticación no configurada")
        repository = auth_repository_provider()
        user = repository.authenticate_user(request.email)
        if not user or not user.get("active") or not verify_password(request.password, user.get("password_hash") or ""):
            raise HTTPException(status_code=401, detail="Credenciales incorrectas")
        repository.record_user_login(int(user["id"]), True, request.email)
        token = create_token(user, auth_secret, auth_token_ttl_seconds)
        user.pop("password_hash", None)
        return {"access_token": token, "token_type": "bearer", "expires_in": auth_token_ttl_seconds, "user": user}

    @router.get("/auth/me", tags=["auth"])
    def auth_me(request: Request):
        authorization = request.headers.get("Authorization", "")
        token = authorization[7:] if authorization.lower().startswith("bearer ") else ""
        claims = decode_token(token, auth_secret)
        if not claims or auth_repository_provider is None:
            raise HTTPException(status_code=401, detail="Sesión no válida")
        user = auth_repository_provider().get_user(int(claims["sub"]))
        if not user or not user.get("active"):
            raise HTTPException(status_code=401, detail="Usuario no disponible")
        user.pop("password_hash", None)
        return user

    @router.get("/integrations", tags=["integrations"])
    def integrations(current: Annotated[IntegrationService, Depends(integration_service)]):
        return current.list()

    @router.get("/integrations/health", tags=["integrations"])
    def integrations_health(current: Annotated[IntegrationService, Depends(integration_service)]):
        return current.health()

    @router.post("/integrations/{code}/sync", tags=["integrations"])
    def integration_sync(code: str, request: IntegrationSyncRequest, current: Annotated[IntegrationService, Depends(integration_service)]):
        try:
            return current.sync(code, request.direction)
        except ValueError as error:
            raise HTTPException(status_code=404, detail=str(error)) from error

    @router.get("/health", response_model=HealthResponse, tags=["system"])
    def health() -> HealthResponse:
        return HealthResponse(status="ok", service="smart-warehouse-api", environment=environment)

    @router.get("/dashboard", tags=["dashboard"])
    def dashboard(current: Annotated[WarehouseService, Depends(service)]):
        return current.dashboard()

    @router.get("/orders", tags=["orders"])
    def orders(current: Annotated[WarehouseService, Depends(service)], status: str | None = Query(default=None)):
        return current.orders(status=status)

    @router.get("/orders/{external_id}", tags=["orders"])
    def order(external_id: str, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.order(external_id)
        except OrderNotFoundError as error:
            raise HTTPException(status_code=404, detail=f"Order {external_id} not found") from error

    @router.post("/orders/{external_id}/status", tags=["orders"])
    def transition_order_status(external_id: str, request: OrderStatusRequest, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.transition_order_status(external_id, request.status, request.reason)
        except ValueError as error:
            raise HTTPException(status_code=409, detail=str(error)) from error

    @router.get("/orders/{external_id}/status-history", tags=["orders"])
    def order_status_history(external_id: str, current: Annotated[WarehouseService, Depends(service)]):
        return current.order_status_history(external_id)

    @router.get("/orders/{external_id}/decisions", tags=["orders"])
    def decisions(external_id: str, current: Annotated[WarehouseService, Depends(service)]):
        return current.decisions(external_id)

    @router.get("/ai/anomalies", tags=["ai"])
    def ai_anomalies(current: Annotated[AIService, Depends(ai_service)]):
        return current.anomalies(persist=False)

    @router.post("/ai/anomalies/run", tags=["ai"])
    def run_ai_anomalies(current: Annotated[AIService, Depends(ai_service)]):
        return current.anomalies(persist=True)

    @router.get("/ai/demand", tags=["ai"])
    def ai_demand(current: Annotated[AIService, Depends(ai_service)], sku: str | None = None):
        return current.demand(sku=sku)

    @router.get("/ai/suppliers/compare", tags=["ai"])
    def ai_supplier_comparison(current: Annotated[AIService, Depends(ai_service)]):
        return current.suppliers()

    @router.get("/ai/suggestions", tags=["ai"])
    def ai_suggestions(current: Annotated[AIService, Depends(ai_service)]):
        return current.suggestions(persist=False)

    @router.post("/ai/suggestions/run", tags=["ai"])
    def run_ai_suggestions(current: Annotated[AIService, Depends(ai_service)]):
        return current.suggestions(persist=True)

    @router.post("/ai/chat", tags=["ai"])
    def ai_chat(request: AIChatRequest, current: Annotated[AIService, Depends(ai_service)]):
        return current.chat(request.message)

    @router.post("/orders/{external_id}/validate", response_model=ValidationResponse, tags=["orders"])
    def validate_order(external_id: str, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.validate(external_id)
        except OrderNotFoundError as error:
            raise HTTPException(status_code=404, detail=f"Order {external_id} not found") from error

    @router.post("/orders", status_code=201, tags=["orders"])
    def create_order(payload: CreateOrderRequest, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.create_order(payload.model_dump())
        except ValueError as error:
            raise HTTPException(status_code=422, detail=str(error)) from error

    @router.post("/imports/orders", tags=["imports"])
    async def import_orders(file: UploadFile = File(...), confirm: bool = Form(False), current: WarehouseService = Depends(service)):
        filename = (file.filename or "").lower()
        normalized, parse_errors = _parse_order_import(filename, await file.read())
        preview = current.preview_import_orders(normalized)
        errors = parse_errors + preview["errors"]
        error_rows = {item["row"] for item in errors}
        valid_rows = [row for row in normalized if row.get("_row") not in error_rows]
        if not confirm:
            return {"mode": "preview", "preview_count": len(normalized), "valid_count": len(valid_rows), "error_count": len(errors), "rows": normalized, "errors": errors}
        result = current.import_orders(valid_rows)
        return {"mode": "imported", "preview_count": len(normalized), **result, "errors": errors + result.get("errors", []), "error_count": len(errors) + result.get("error_count", 0)}

    @router.get("/stock", tags=["stock"])
    def stock(current: Annotated[WarehouseService, Depends(service)]):
        return current.stock()

    @router.get("/stock/movements", tags=["stock"])
    def stock_movements(current: Annotated[WarehouseService, Depends(service)], sku: str | None = None, limit: int = Query(default=100, ge=1, le=500)):
        return current.stock_movements(sku=sku, limit=limit)

    @router.post("/stock/movements", status_code=201, tags=["stock"])
    def create_stock_movement(request: StockMovementRequest, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.create_stock_movement(request.model_dump())
        except ValueError as error:
            raise HTTPException(status_code=422, detail=str(error)) from error

    @router.get("/suppliers", tags=["suppliers"])
    def suppliers(current: Annotated[WarehouseService, Depends(service)]):
        return current.suppliers()

    @router.get("/documents/invoices", tags=["documents"])
    def invoices(current: Annotated[WarehouseService, Depends(service)]):
        return current.invoices()

    @router.post("/documents/invoices", status_code=201, tags=["documents"])
    async def upload_invoice(file: UploadFile = File(...), current: WarehouseService = Depends(service)):
        filename = file.filename or "invoice.pdf"
        if not filename.lower().endswith(".pdf") and file.content_type != "application/pdf":
            raise HTTPException(status_code=415, detail="Solo se admiten facturas PDF")
        content = await file.read()
        if not content:
            raise HTTPException(status_code=422, detail="La factura está vacía")
        extracted = _extract_invoice(content)
        try:
            return current.save_invoice_document(filename, file.content_type or "application/pdf", content, extracted)
        except ValueError as error:
            raise HTTPException(status_code=422, detail=str(error)) from error

    @router.post("/documents/invoices/{invoice_number}/reconcile", tags=["documents"])
    def reconcile_invoice(invoice_number: str, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.reconcile_invoice(invoice_number)
        except ValueError as error:
            raise HTTPException(status_code=404, detail=str(error)) from error

    @router.get("/documents/invoices/accounting-export", tags=["documents"])
    def accounting_export(format: str = Query(default="csv", pattern="^(csv|xlsx)$"), current: WarehouseService = Depends(service)):
        payload, media_type, filename = _accounting_file(current.accounting_export_rows(), format)
        return StreamingResponse(io.BytesIO(payload), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'})

    @router.post("/documents/invoices/{invoice_number}/accounting-export", tags=["documents"])
    def export_invoice(invoice_number: str, request: AccountingExportRequest, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.mark_invoice_exported(invoice_number, "rest", request.target_system)
        except ValueError as error:
            raise HTTPException(status_code=409, detail=str(error)) from error

    @router.get("/documents/procedures", tags=["documents"])
    def procedures(current: Annotated[WarehouseService, Depends(service)]):
        return current.procedures()

    @router.get("/orders/{external_id}/procedures", tags=["documents"])
    def order_procedures(external_id: str, current: Annotated[WarehouseService, Depends(service)]):
        return current.procedures(external_id=external_id)

    @router.post("/orders/{external_id}/procedures/{procedure_code}/documents", status_code=201, tags=["documents"])
    async def upload_procedure_document(external_id: str, procedure_code: str, file: UploadFile = File(...), current: WarehouseService = Depends(service)):
        filename = file.filename or "procedure-document"
        content = await file.read()
        if not content:
            raise HTTPException(status_code=422, detail="El documento está vacío")
        try:
            return current.save_procedure_document(external_id, procedure_code, filename, file.content_type or "application/octet-stream", content)
        except ValueError as error:
            raise HTTPException(status_code=422, detail=str(error)) from error

    @router.get("/receipts", tags=["receiving"])
    def receipts(current: Annotated[WarehouseService, Depends(service)]):
        return current.receipts()

    @router.post("/receipts", status_code=201, tags=["receiving"])
    def create_receipt(request: CreateReceiptRequest, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.create_receipt(request.model_dump())
        except ValueError as error:
            raise HTTPException(status_code=422, detail=str(error)) from error

    @router.get("/events", tags=["events"])
    def events(current: Annotated[WarehouseService, Depends(service)], limit: int = Query(default=50, ge=1, le=200), severity: str | None = Query(default=None, pattern="^(info|warning|critical)$"), event_type: str | None = None, aggregate_type: str | None = None):
        return current.events(limit=limit, severity=severity, event_type=event_type, aggregate_type=aggregate_type)

    @router.get("/events/export", tags=["events"])
    def export_events(format: str = Query(default="csv", pattern="^(csv|xlsx)$"), severity: str | None = Query(default=None, pattern="^(info|warning|critical)$"), event_type: str | None = None, aggregate_type: str | None = None, current: WarehouseService = Depends(service)):
        payload, media_type, filename = _events_file(current.events(limit=200, severity=severity, event_type=event_type, aggregate_type=aggregate_type), format)
        return StreamingResponse(io.BytesIO(payload), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'})

    @router.get("/events/{event_id}", tags=["events"])
    def event_detail(event_id: int, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.event(event_id)
        except ValueError as error:
            raise HTTPException(status_code=404, detail=str(error)) from error

    @router.get("/alerts", tags=["events"])
    def alerts(current: Annotated[WarehouseService, Depends(service)], limit: int = Query(default=50, ge=1, le=200), severity: str | None = Query(default=None, pattern="^(info|warning|critical)$"), event_type: str | None = None, status: str = Query(default="unread", pattern="^(unread|read|all)$")):
        return current.alerts(limit=limit, severity=severity, event_type=event_type, status=status)

    @router.post("/alerts/read-all", tags=["events"])
    def read_all_alerts(current: Annotated[WarehouseService, Depends(service)]):
        return {"marked_read": current.mark_all_alerts_read()}

    @router.post("/alerts/{alert_id}/read", tags=["events"])
    def read_alert(alert_id: int, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.mark_alert_read(alert_id)
        except ValueError as error:
            raise HTTPException(status_code=404, detail=str(error)) from error

    @router.get("/alert-rules", tags=["events"])
    def alert_rules(current: Annotated[WarehouseService, Depends(service)]):
        return current.alert_rules()

    @router.post("/alert-rules", status_code=201, tags=["events"])
    def create_alert_rule(request: AlertRuleRequest, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.create_alert_rule(request.model_dump())
        except ValueError as error:
            raise HTTPException(status_code=409, detail=str(error)) from error

    @router.patch("/alert-rules/{rule_id}", tags=["events"])
    def update_alert_rule(rule_id: int, request: AlertRuleRequest, current: Annotated[WarehouseService, Depends(service)]):
        try:
            return current.update_alert_rule(rule_id, request.model_dump())
        except ValueError as error:
            raise HTTPException(status_code=404, detail=str(error)) from error

    return router
